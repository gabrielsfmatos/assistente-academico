"""Importa a planilha "Gestão Faculdade" para o banco do sistema.

Uso (na raiz do projeto, com o ambiente virtual ativo):
    python -m scripts.import_xlsx "data\\Gestao_Faculdade.xlsx" --reset

Além de importar, o script audita as fórmulas da planilha e gera
data/relatorio_importacao.md com as inconsistências encontradas
(material do "antes" para o AI Show Case).
"""

import argparse
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from sqlmodel import select

from app.db import DB_PATH, get_engine, get_session, init_db
from app.models import Assessment, ClassSession, Course, CourseStatus, ScheduleSlot
from app.services.attendance import compute_summary
from app.services.grades import APPROVAL_GRADE, compute_cr, cr_by_term

# Siglas usadas na grade de horários -> nome da disciplina.
CODE_TO_NAME = {
    "ACD": "Algoritmos para Ciência de Dados",
    "LPP": "Linguagens e Paradigmas de Programação",
    "EI": "Empreendedorismo e Inovação",
    "CDD": "Ciência de Dados",
    "AP": "Aprendizagem Profunda",
}
# Curso atribuído às disciplinas que a fórmula do CR exclui (cursadas em outro curso).
OTHER_PROGRAM = "EP"  # Engenharia de Produção

WEEKDAY_COLUMNS = {"B": 0, "C": 1, "D": 2, "E": 3, "F": 4}  # 2ª a 6ª


def normalize(name: str) -> str:
    """'Empreendedorismo e Inovação* ' -> 'empreendedorismo e inovação'"""
    return name.replace("*", "").strip().casefold()


def parse_attendance_cell(value) -> dict | None:
    """Converte uma célula do Faltômetro nos campos de ClassSession.

    Retorna None quando não há aula no dia (célula vazia ou só espaços).
    Prova e teste viram presença (regra: tudo é presença ou falta).
    A chave "exam" não é campo da aula: indica ao importador que aquele
    dia tem prova, que vira um registro em Assessment (fonte única da data).
    """
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if value is True:
        return {"attended": True}
    if value is False:
        return {"attended": False}
    text = str(value).strip().upper()
    if text == "S/ AULA":
        return {"attended": True, "cancelled": True}
    if text == "PROVA":
        return {"attended": True, "exam": True}
    if text == "TESTE":
        return {"attended": True, "has_quiz": True}
    raise ValueError(f"Valor não reconhecido no Faltômetro: {value!r}")


def parse_time_range(label: str) -> tuple[str, str]:
    """'16h - 18h' -> ('16:00', '18:00')"""
    hours = re.findall(r"(\d{1,2})h", label)
    if len(hours) != 2:
        raise ValueError(f"Horário não reconhecido: {label!r}")
    return f"{int(hours[0]):02d}:00", f"{int(hours[1]):02d}:00"


def as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


# ------------------------------------------------------------------ importação


def import_cr_sheet(ws, report: list[str]) -> list[Course]:
    courses: list[Course] = []
    row_of: dict[int, int] = {}  # id(course) -> linha na planilha
    current_term = None
    excluded_rows: set[int] = set()

    for row in range(3, ws.max_row + 1):
        label = ws.cell(row, 1).value
        if isinstance(label, str) and label.strip().startswith("CR Acumulado"):
            formula = str(ws.cell(row, 2).value or "")
            # Ex.: "- C17*D17" na fórmula do CR acumulado = linha 17 excluída
            excluded_rows = {int(r) for r in re.findall(r"-\s*C(\d+)\s*\*\s*D\1", formula)}
            break
        if isinstance(label, str) and "Período" in label:
            current_term = int(re.match(r"\s*(\d+)", label).group(1))
        name = ws.cell(row, 2).value
        if not name or current_term is None:
            continue
        grade = ws.cell(row, 3).value
        credits = ws.cell(row, 4).value
        final_grade = float(grade) if grade is not None else None
        if final_grade is None:
            status = CourseStatus.IN_PROGRESS
        elif final_grade >= APPROVAL_GRADE:
            status = CourseStatus.PASSED
        else:
            status = CourseStatus.FAILED
        course = Course(
            name=str(name).strip(),
            term=current_term,
            credits=float(credits or 0),
            final_grade=final_grade,
            status=status,
        )
        row_of[id(course)] = row
        courses.append(course)

    for c in courses:
        row = row_of[id(c)]
        if row in excluded_rows:
            c.program = OTHER_PROGRAM
            report.append(
                f"- **{c.name}** é excluída do CR por uma conta escondida na fórmula "
                f"(`- C{row}*D{row}`). No sistema, virou um atributo explícito: "
                f"curso = {OTHER_PROGRAM}."
            )
    return courses


def import_attendance_sheet(ws, by_name: dict, report: list[str]):
    # Linha 3: datas. Guardamos coluna -> data.
    date_cols = {
        cell.column: as_date(cell.value) for cell in ws[3] if as_date(cell.value) is not None
    }
    last_date_col = max(date_cols)
    sessions_by_course: dict[str, list[ClassSession]] = {}
    exam_days_by_course: dict[str, list[date]] = {}
    prova_rule: dict[str, bool] = {}

    for row in range(4, ws.max_row + 1):
        raw_name = ws.cell(row, 1).value
        if not raw_name or str(raw_name).strip().startswith("*"):
            continue
        key = normalize(str(raw_name))
        course = by_name.get(key)
        if course is None:
            report.append(f"- Disciplina do Faltômetro sem correspondência no histórico: {raw_name!r}")
            continue
        course.requires_attendance = "*" not in str(raw_name)

        sessions = []
        blank_strings = []
        for col, day in date_cols.items():
            value = ws.cell(row, col).value
            if isinstance(value, str) and value != "" and not value.strip():
                blank_strings.append(day)
            fields = parse_attendance_cell(value)
            if fields is not None:
                if fields.pop("exam", False):
                    exam_days_by_course.setdefault(key, []).append(day)
                sessions.append(ClassSession(day=day, **fields))
        sessions_by_course[key] = sessions

        if blank_strings:
            dias = ", ".join(f"{d:%d/%m}" for d in blank_strings)
            report.append(f"- **{course.name}**: células com espaço em branco em {dias}.")

        # ---- auditoria das fórmulas de Faltas (CN) e Frequência (CO)
        absences_formula = str(ws.cell(row, last_date_col + 1).value or "")
        freq_formula = str(ws.cell(row, last_date_col + 2).value or "")

        if re.search(r"\*\s*2\b", absences_formula):
            report.append(f"- **{course.name}**: as faltas são multiplicadas por 2 (`{absences_formula}`).")

        prova_rule[course.name] = "PROVA" in freq_formula

        m = re.search(r"B\d+:([A-Z]+)\d+", absences_formula)
        if m:
            end_col = column_index_from_string(m.group(1))
            ignored = [d for c, d in date_cols.items() if c > end_col]
            ignored_classes = [s for s in sessions if s.day in ignored and not s.cancelled]
            if ignored_classes:
                report.append(
                    f"- **{course.name}**: a fórmula vai só até a coluna {m.group(1)} e ignora "
                    f"{len(ignored_classes)} aula(s) a partir de {min(ignored):%d/%m}."
                )

        quizzes = [s for s in sessions if s.has_quiz]
        if quizzes and "TESTE" not in freq_formula:
            dias = ", ".join(f"{s.day:%d/%m}" for s in quizzes)
            report.append(f"- **{course.name}**: o dia de TESTE ({dias}) não entra em nenhuma conta.")

    counts_exam = [n for n, v in prova_rule.items() if v]
    ignores_exam = [n for n, v in prova_rule.items() if not v]
    if counts_exam and ignores_exam:
        report.append(
            "- Regra de PROVA inconsistente: conta como presença em "
            f"{', '.join(counts_exam)}, mas não em {', '.join(ignores_exam)}."
        )
    return sessions_by_course, exam_days_by_course, last_date_col


def import_schedule_and_exams(ws, by_name: dict, report: list[str]):
    slots = []
    for row in range(3, ws.max_row + 1):
        label = ws.cell(row, 1).value
        if not label or "h" not in str(label):
            continue
        start, end = parse_time_range(str(label))
        for col_letter, weekday in WEEKDAY_COLUMNS.items():
            code = ws[f"{col_letter}{row}"].value
            if not code or str(code).strip() in {"-", ""}:
                continue
            code = str(code).strip()
            name = CODE_TO_NAME.get(code)
            course = by_name.get(normalize(name)) if name else None
            if course is None:
                report.append(f"- Sigla da grade sem disciplina correspondente: {code!r}")
                continue
            course.code = code
            slots.append((course, ScheduleSlot(weekday=weekday, start_time=start, end_time=end)))

    exams = []
    for row in range(3, ws.max_row + 1):
        raw_name = ws.cell(row, 8).value  # coluna H
        if not raw_name:
            continue
        course = by_name.get(normalize(str(raw_name)))
        if course is None:
            report.append(f"- Disciplina da tabela de provas sem correspondência: {raw_name!r}")
            continue
        for label, date_col, grade_col in (("P1", 9, 10), ("P2", 11, 12)):
            day = as_date(ws.cell(row, date_col).value)
            grade = ws.cell(row, grade_col).value
            if day or grade is not None:
                exams.append((course, Assessment(name=label, day=day, grade=grade)))
        if not any(ws.cell(row, c).value for c in (9, 10, 11, 12)):
            report.append(f"- **{course.name}** aparece na tabela de provas, mas sem datas nem notas.")
    return slots, exams


# ------------------------------------------------------------------ principal


def run(xlsx_path: Path, reset: bool) -> None:
    if reset and DB_PATH.exists():
        DB_PATH.unlink()

    engine = get_engine()
    init_db(engine)
    report: list[str] = []

    wb = load_workbook(xlsx_path)  # fórmulas
    wb_values = load_workbook(xlsx_path, data_only=True)  # valores calculados pelo Excel/Sheets

    with get_session(engine) as db:
        if db.exec(select(Course)).first() is not None:
            sys.exit("O banco já tem dados. Rode de novo com --reset para recriar do zero.")

        courses = import_cr_sheet(wb["Calculadora de CR"], report)
        by_name = {normalize(c.name): c for c in courses}
        sessions_by_course, exam_days_by_course, last_col = import_attendance_sheet(
            wb["Faltômetro"], by_name, report
        )
        slots, exams = import_schedule_and_exams(wb["Horário e Datas"], by_name, report)

        for c in courses:
            db.add(c)
        db.commit()
        for c in courses:
            db.refresh(c)
            for s in sessions_by_course.get(normalize(c.name), []):
                s.course_id = c.id
                db.add(s)
        for course, slot in slots:
            slot.course_id = course.id
            db.add(slot)
        for course, exam in exams:
            exam.course_id = course.id
            db.add(exam)
        # Dias marcados como PROVA no Faltômetro viram avaliações, se a
        # tabela de provas ainda não tiver uma avaliação nesse dia.
        for c in courses:
            known = {e.day for course, e in exams if course is c}
            for i, day in enumerate(exam_days_by_course.get(normalize(c.name), []), start=1):
                if day not in known:
                    db.add(Assessment(course_id=c.id, name=f"Prova {day:%d/%m}", day=day))
                    report.append(
                        f"- **{c.name}**: PROVA em {day:%d/%m} no Faltômetro, mas fora da "
                        "tabela de provas. Cadastrada como avaliação."
                    )
        db.commit()

        # ---- conferência: aulas em dias que não batem com a grade
        for c in courses:
            weekdays = {s.weekday for s in c.schedule}
            if not weekdays:
                continue
            odd = [s.day for s in c.sessions if s.day.weekday() not in weekdays]
            if odd:
                report.append(
                    f"- **{c.name}**: aulas em dias fora da grade: "
                    + ", ".join(f"{d:%d/%m}" for d in odd)
                )

        # ---- comparação planilha x sistema
        ws_val = wb_values["Faltômetro"]
        comparison = [
            "| Disciplina | Faltas (planilha) | Faltas (sistema) | Freq. projetada (planilha) "
            "| Freq. projetada (sistema) | Faltas permitidas | Restantes |",
            "|---|---|---|---|---|---|---|",
        ]
        today = date.today()
        for row in range(4, ws_val.max_row + 1):
            raw = ws_val.cell(row, 1).value
            course = by_name.get(normalize(str(raw))) if raw else None
            if course is None:
                continue
            sheet_abs = ws_val.cell(row, last_col + 1).value
            sheet_freq = ws_val.cell(row, last_col + 2).value
            s = compute_summary(course.sessions, today, course.name, course.requires_attendance)
            fmt = lambda v: f"{v:.1%}" if isinstance(v, (int, float)) else "—"  # noqa: E731
            comparison.append(
                f"| {course.name} | {sheet_abs} | {s.absences} | {fmt(sheet_freq)} | "
                f"{fmt(s.projected_frequency)} | {s.max_absences} | {s.remaining_absences} |"
            )

        cr_sheet = wb_values["Calculadora de CR"]
        sheet_cr = next(
            (cr_sheet.cell(r, 2).value for r in range(1, cr_sheet.max_row + 1)
             if str(cr_sheet.cell(r, 1).value or "").startswith("CR Acumulado")),
            None,
        )
        all_courses = db.exec(select(Course)).all()
        system_cr = compute_cr(all_courses)
        terms = cr_by_term(all_courses)

    lines = [
        "# Relatório de importação",
        "",
        f"Arquivo: `{xlsx_path.name}` · gerado em {datetime.now():%d/%m/%Y %H:%M}",
        "",
        "## Inconsistências encontradas na planilha",
        "",
        *(report or ["- Nenhuma."]),
        "",
        "## Frequência: planilha x sistema",
        "",
        *comparison,
        "",
        "## CR",
        "",
        f"- CR acumulado na planilha: **{sheet_cr}**",
        f"- CR acumulado calculado pelo sistema: **{system_cr}**",
        "- CR por período: " + ", ".join(f"{t}º = {v}" for t, v in terms.items() if v is not None),
        "",
    ]
    out = Path("data") / "relatorio_importacao.md"
    out.parent.mkdir(exist_ok=True)
    out.write_text("\n".join(lines), encoding="utf-8")
    print("\n".join(lines))
    print(f"\nBanco criado em: {DB_PATH}\nRelatório salvo em: {out}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Importa a planilha Gestão Faculdade.")
    parser.add_argument("xlsx", type=Path, help="caminho do arquivo .xlsx")
    parser.add_argument("--reset", action="store_true", help="apaga o banco e importa do zero")
    args = parser.parse_args()
    if not args.xlsx.exists():
        sys.exit(f"Arquivo não encontrado: {args.xlsx}")
    run(args.xlsx, args.reset)


if __name__ == "__main__":
    main()
